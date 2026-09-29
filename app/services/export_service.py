"""Export the job tracker (jobs joined to their applications) to Excel and CSV.

Job text comes from untrusted sources, so every text cell is neutralised against spreadsheet
formula injection and stripped of control characters. Only tracker fields are exported: no
approval tokens, form answers, screenshots or resume contents.
"""

from __future__ import annotations

import csv
import os
import tempfile
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Literal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import Settings
from app.database.base import utcnow
from app.database.tables import ApplicationRecord, JobRecord
from app.utils.file_utils import ensure_directories
from app.utils.logging_config import audit

ExportFormat = Literal["xlsx", "csv"]
FORMATS: tuple[ExportFormat, ...] = ("xlsx", "csv")
BASE_NAME = "applications"
MAX_CELL_CHARS = 32000
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")

Cell = str | float | None

HEADERS: tuple[str, ...] = (
    "Application ID",
    "Job ID",
    "Company",
    "Job Title",
    "Location",
    "Employment Type",
    "Workplace Type",
    "Source",
    "Portal",
    "Job URL",
    "Date Discovered (UTC)",
    "Match Score",
    "Recommendation",
    "Matching Explanation",
    "Missing Requirements",
    "Tailored Resume File",
    "Status",
    "Application Started (UTC)",
    "Awaiting Approval (UTC)",
    "Approved (UTC)",
    "Submitted (UTC)",
    "Confirmation Number",
    "Failure Reason",
    "Notes",
    "Last Updated (UTC)",
)


class ExportError(RuntimeError):
    pass


class ExportFileLockedError(ExportError):
    """The target file could not be replaced, usually because it is open in Excel."""


@dataclass(frozen=True)
class ExportResult:
    rows: int
    generated_at: datetime
    files: dict[ExportFormat, Path]


def clean_text(value: str) -> str:
    """Drop control characters, cap the length and defuse leading formula characters."""
    text = ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_CHARS]
    return f"'{text}" if text.startswith(FORMULA_PREFIXES) else text


def _text(value: str | None) -> str:
    return clean_text(value) if value else ""


def _moment(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value is not None else ""


def build_row(job: JobRecord, application: ApplicationRecord | None) -> list[Cell]:
    """One tracker row; the application record wins over the job once one exists."""
    app = application
    missing = (app.missing_requirements if app else None) or job.missing_requirements or []
    score = app.match_score if app and app.match_score is not None else job.match_score
    recommendation = (app.recommendation if app else None) or job.recommendation
    explanation = (app.matching_explanation if app else None) or job.matching_explanation
    resume = app.tailored_resume_path if app else None
    return [
        app.application_id if app else "",
        job.job_id,
        _text(app.company if app else job.company),
        _text(app.job_title if app else job.title),
        _text(job.location),
        _text(job.employment_type),
        _text(job.workplace_type),
        _text(job.source),
        _text(app.portal if app else None),
        _text(job.job_url),
        _moment(job.date_discovered),
        score,
        _text(recommendation),
        _text(explanation),
        _text("; ".join(missing)),
        _text(Path(resume).name if resume else None),
        app.application_status if app else job.status,
        _moment(app.application_started_at if app else None),
        _moment(app.awaiting_approval_at if app else None),
        _moment(app.approved_at if app else None),
        _moment(app.submitted_at if app else None),
        _text(app.confirmation_number if app else None),
        _text(app.failure_reason if app else None),
        _text(app.notes if app else None),
        _moment(app.updated_at if app else job.updated_at),
    ]


def _publish(target: Path, write: Callable[[Path], None]) -> None:
    """Write to a temporary file beside the target, then replace it atomically."""
    fd, tmp_name = tempfile.mkstemp(dir=target.parent, suffix=".tmp")
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        write(tmp)
        os.replace(tmp, target)
    except PermissionError as exc:
        raise ExportFileLockedError(
            f"Cannot replace {target.name}; close it in Excel and try again."
        ) from exc
    finally:
        tmp.unlink(missing_ok=True)


def _write_csv(path: Path, rows: list[list[Cell]]) -> None:
    # utf-8-sig so Excel detects the encoding when the CSV is opened directly.
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for row in rows:
            writer.writerow(["" if cell is None else cell for cell in row])


def _write_xlsx(path: Path, rows: list[list[Cell]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise ExportError("Cannot create worksheet")
    sheet.title = "Applications"
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for line in sheet.iter_rows(min_row=2):
        for cell in line:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(HEADERS, start=1):
        longest = max((len(str(row[index - 1] or "")) for row in rows), default=0)
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(header), longest) + 2, 60
        )
    workbook.save(path)


class TrackerExporter:
    def __init__(self, session: Session, settings: Settings) -> None:
        self.session = session
        self.settings = settings

    def rows(self) -> list[list[Cell]]:
        query = (
            select(JobRecord, ApplicationRecord)
            .outerjoin(ApplicationRecord, ApplicationRecord.job_id == JobRecord.job_id)
            .order_by(JobRecord.id)
        )
        return [build_row(job, application) for job, application in self.session.execute(query)]

    def export(self, formats: tuple[ExportFormat, ...] = FORMATS) -> ExportResult:
        rows = self.rows()
        directory = self.settings.exports_dir
        ensure_directories([directory])
        files: dict[ExportFormat, Path] = {}
        for fmt in dict.fromkeys(formats):
            target = directory / f"{BASE_NAME}.{fmt}"
            if fmt == "xlsx":
                _publish(target, lambda tmp: _write_xlsx(tmp, rows))
            else:
                _publish(target, lambda tmp: _write_csv(tmp, rows))
            files[fmt] = target
        audit("tracker_export", "ok", rows=len(rows), formats=",".join(files))
        return ExportResult(rows=len(rows), generated_at=utcnow(), files=files)
