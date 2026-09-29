"""Final end-to-end run: a real `run.py serve` process driven over HTTP, synthetic site only.

Nothing here talks to a real job portal. The only page that can be submitted is the local mock.
"""

from __future__ import annotations

import csv
import os
import socket
import sqlite3
import subprocess
import sys
import time
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import httpx
import pytest
import yaml
from openpyxl import load_workbook

from app.config import PROJECT_ROOT, Settings
from app.models.candidate import CandidateProfile
from app.services.export_service import HEADERS

from .test_browser_forms import MockSite, browser_channel, browser_settings, site  # noqa: F401
from .test_matching_api import WEAK, job
from .test_resume_tailoring import MASTER, build_master

API_KEY = "e2e-test-key-0123456789abcdef0123456789abcdef"
SETTING_NAMES = {name.upper() for name in Settings.model_fields}


@dataclass
class LiveServer:
    root: Path
    url: str
    http: httpx.Client
    process: subprocess.Popen[bytes]

    def rows(self, sql: str, *args: Any) -> list[tuple[Any, ...]]:
        with sqlite3.connect(self.root / "data" / "job_tracker.db") as connection:
            return connection.execute(sql, args).fetchall()


def _free_port() -> int:
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        return int(probe.getsockname()[1])


def _environment(root: Path, channel: str, **extra: str) -> dict[str, str]:
    env = {k: v for k, v in os.environ.items() if k.upper() not in SETTING_NAMES}
    env.update(
        PROJECT_ROOT=str(root),
        CONFIG_DIRECTORY="config",
        PLAYWRIGHT_HEADLESS="true",
        PLAYWRIGHT_SLOW_MO="0",
        PLAYWRIGHT_TIMEOUT="10000",
        PLAYWRIGHT_BROWSER_CHANNEL=channel,
        API_KEY=API_KEY,
        AI_ENABLED="false",
    )
    env.update(extra)
    return env


def _serve(root: Path, channel: str, port: int, **extra: str) -> subprocess.Popen[bytes]:
    log = (root / "server.out").open("wb")
    return subprocess.Popen(
        [sys.executable, str(PROJECT_ROOT / "run.py"), "serve", "--port", str(port)],
        cwd=PROJECT_ROOT,
        env=_environment(root, channel, **extra),
        stdout=log,
        stderr=subprocess.STDOUT,
    )


@pytest.fixture
def live(browser_settings: Settings, browser_channel: str) -> Iterator[LiveServer]:  # noqa: F811
    root = browser_settings.project_root
    path = browser_settings.config_dir / "candidate_profile.yaml"
    profile = CandidateProfile.model_validate(yaml.safe_load(path.read_text(encoding="utf-8")))
    build_master(root / MASTER, profile)
    port = _free_port()
    process = _serve(root, browser_channel, port)
    url = f"http://127.0.0.1:{port}"
    http = httpx.Client(base_url=url, headers={"X-API-Key": API_KEY}, timeout=120, trust_env=False)
    try:
        deadline = time.monotonic() + 60
        while True:
            if process.poll() is not None:
                pytest.fail((root / "server.out").read_text(errors="replace")[-2000:])
            try:
                if http.get("/health").status_code == 200:
                    break
            except httpx.TransportError:
                pass
            if time.monotonic() > deadline:
                pytest.fail("server did not start")
            time.sleep(0.3)
        yield LiveServer(root, url, http, process)
    finally:
        http.close()
        process.terminate()
        try:
            process.wait(timeout=15)
        except subprocess.TimeoutExpired:
            process.kill()


def test_full_pipeline_over_http_needs_the_typed_approval(
    live: LiveServer,
    site: MockSite,  # noqa: F811
) -> None:
    http = live.http

    health = http.get("/health").json()
    assert health["status"] == "ok" and health["startup_ok"] is True
    assert health["safety"]["manual_approval_required"] is True
    assert health["safety"]["automatic_submission_enabled"] is False
    assert health["safety"]["ai_enabled"] is False
    assert health["safety"]["violations"] == []
    assert httpx.get(f"{live.url}/jobs", trust_env=False).status_code == 401

    strong = job("Acme Networks", job_url=f"{site.base}/mock_application.html")
    weak = job("Weak Co", WEAK)
    imported = http.post("/jobs/import/json", json={"jobs": [strong, weak, strong]}).json()
    assert imported["created"] == 2 and imported["duplicates"] == 1
    strong_id, weak_id = (item["job_id"] for item in imported["results"][:2])

    assert http.post(f"/matching/analyze/{strong_id}").json()["result"]["recommendation"] == "APPLY"
    assert http.post(f"/matching/analyze/{weak_id}").json()["result"]["recommendation"] != "APPLY"
    assert http.post(f"/resumes/tailor/{weak_id}").status_code == 409

    tailored = http.post(f"/resumes/tailor/{strong_id}")
    assert tailored.status_code == 200, tailored.text
    assert tailored.json()["validation"]["passed"] is True
    application_id = tailored.json()["application_id"]
    assert (live.root / tailored.json()["path"]).is_file()

    form = http.post(f"/applications/{application_id}/form/start", json={})
    assert form.status_code == 200, form.text
    assert form.json()["outcome"] == "READY_FOR_REVIEW"
    assert form.json()["application_status"] == "AWAITING_APPROVAL"
    assert site.posts == []

    # Nothing can submit without a fresh, matching approval.
    forged = http.post(
        f"/applications/{application_id}/approval/confirm",
        json={"approval_token": "x" * 43, "approval_text": f"SUBMIT {application_id}"},
    )
    assert forged.status_code == 403
    assert site.posts == []

    challenge = http.post(f"/applications/{application_id}/approval/request").json()
    text = f"SUBMIT {application_id}"
    assert challenge["required_text"] == text
    for wrong in ("SUBMIT", f"submit {application_id}"):
        response = http.post(
            f"/applications/{application_id}/approval/confirm",
            json={"approval_token": challenge["approval_token"], "approval_text": wrong},
        )
        assert response.status_code == 403
    assert site.posts == []

    done = http.post(
        f"/applications/{application_id}/approval/confirm",
        json={"approval_token": challenge["approval_token"], "approval_text": text},
    )
    assert done.status_code == 200, done.text
    assert done.json()["submitted"] is True
    assert done.json()["confirmation_number"] == "SYN-12345"
    assert site.posts == ["/submit"]

    replay = http.post(
        f"/applications/{application_id}/approval/confirm",
        json={"approval_token": challenge["approval_token"], "approval_text": text},
    )
    assert replay.status_code == 403
    assert site.posts == ["/submit"]
    assert http.get(f"/applications/{application_id}").json()["application_status"] == "SUBMITTED"

    exported = http.post("/exports/tracker").json()
    assert exported["rows"] == 2
    exports = live.root / "exports"
    with (exports / "applications.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    by_job = {row["Job ID"]: row for row in rows}
    assert by_job[strong_id]["Status"] == "SUBMITTED"
    assert by_job[strong_id]["Application ID"] == application_id
    assert by_job[weak_id]["Application ID"] == ""
    assert by_job[weak_id]["Status"] == http.get(f"/jobs/{weak_id}").json()["status"]
    sheet = load_workbook(exports / "applications.xlsx")["Applications"]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == 3

    history = [
        row[0]
        for row in live.rows(
            "SELECT to_status FROM status_history WHERE application_id = ? ORDER BY id",
            application_id,
        )
    ]
    assert history.index("AWAITING_APPROVAL") < history.index("SUBMITTED") == len(history) - 1
    assert history.count("SUBMITTED") == 1
    attempts = [
        row[0]
        for row in live.rows(
            "SELECT outcome FROM approval_attempts WHERE application_id = ? ORDER BY id",
            application_id,
        )
    ]
    assert [a for a in attempts if a in {"REQUESTED", "APPROVED", "SUBMITTED"}] == [
        "REQUESTED",
        "APPROVED",
        "SUBMITTED",
    ]
    assert attempts.count("REJECTED_TEXT") == 2
    assert attempts.index("REJECTED_TEXT") < attempts.index("APPROVED")
    (approval,) = live.rows("SELECT consumed_at FROM approvals")
    assert approval[0] is not None

    token = challenge["approval_token"]
    with sqlite3.connect(live.root / "data" / "job_tracker.db") as connection:
        assert token not in "\n".join(connection.iterdump())
    for path in [*(live.root / "logs").glob("*.log*"), live.root / "server.out"]:
        assert token not in path.read_text(encoding="utf-8", errors="replace"), path.name


@pytest.mark.parametrize(
    "override",
    [
        {"AUTOMATIC_SUBMISSION_ENABLED": "true"},
        {"MANUAL_APPROVAL_REQUIRED": "false"},
    ],
)
def test_server_refuses_to_start_with_unsafe_flags(
    override: dict[str, str],
    browser_settings: Settings,  # noqa: F811
    browser_channel: str,  # noqa: F811
) -> None:
    process = _serve(browser_settings.project_root, browser_channel, _free_port(), **override)
    assert process.wait(timeout=60) == 1
