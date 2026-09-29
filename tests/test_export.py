from __future__ import annotations

import csv
import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker
from typer.testing import CliRunner

from app.config import PROJECT_ROOT, Settings
from app.database.repositories import ApplicationRepository, JobRepository
from app.database.tables import ApprovalRecord
from app.main import create_app
from app.models.application import ApplicationStatus
from app.services.export_service import HEADERS, TrackerExporter, clean_text
from run import cli

from .test_database import make_job


def _seed(session_factory: sessionmaker[Session]) -> None:
    with session_factory.begin() as session:
        jobs = JobRepository(session)
        applications = ApplicationRepository(session)
        analyzed = jobs.add(make_job("1"))
        analyzed.match_score = 82.5
        analyzed.recommendation = "APPLY"
        analyzed.missing_requirements = ["kubernetes", "go"]
        analyzed.status = "PREPARED"
        application = applications.create_for_job(analyzed, ApplicationStatus.PREPARED)
        application.tailored_resume_path = "C:\\private\\resumes\\generated\\resume-1.docx"
        application.notes = "Synthetic note"
        skipped = jobs.add(make_job("2"))
        skipped.status = "SKIPPED"
        skipped.match_score = 31.0
        session.add(
            ApprovalRecord(
                application_id=application.application_id,
                token_hash="a" * 64,
                expires_at=application.created_at,
            )
        )


def _csv_rows(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def test_export_writes_both_files_with_one_row_per_job(
    settings: Settings, session_factory: sessionmaker[Session]
) -> None:
    _seed(session_factory)
    with session_factory() as session:
        result = TrackerExporter(session, settings).export()
    assert result.rows == 2
    assert {p.name for p in settings.exports_dir.iterdir()} == {
        "applications.csv",
        "applications.xlsx",
    }

    rows = _csv_rows(settings.exports_dir / "applications.csv")
    assert tuple(rows[0]) == HEADERS
    records = [dict(zip(rows[0], row, strict=True)) for row in rows[1:]]
    by_job = {record["Job ID"]: record for record in records}
    prepared = by_job["JOB-1"]
    assert prepared["Application ID"].startswith("APP-")
    assert prepared["Status"] == "PREPARED"
    assert prepared["Match Score"] == "82.5"
    assert prepared["Missing Requirements"] == "kubernetes; go"
    assert prepared["Tailored Resume File"] == "resume-1.docx"
    assert by_job["JOB-2"]["Application ID"] == ""
    assert by_job["JOB-2"]["Status"] == "SKIPPED"

    sheet = load_workbook(settings.exports_dir / "applications.xlsx")["Applications"]
    header = [cell.value for cell in sheet[1]]
    assert tuple(header) == HEADERS
    assert sheet.max_row == 3
    assert sheet.freeze_panes == "A2"
    score_column = HEADERS.index("Match Score") + 1
    assert sheet.cell(2, score_column).value == 82.5


def test_export_leaks_no_secrets_or_paths(
    settings: Settings, session_factory: sessionmaker[Session]
) -> None:
    _seed(session_factory)
    with session_factory() as session:
        TrackerExporter(session, settings).export()
    text = (settings.exports_dir / "applications.csv").read_text(encoding="utf-8-sig")
    assert "a" * 64 not in text
    assert "private" not in text


def test_empty_database_exports_headers_only(
    settings: Settings, session_factory: sessionmaker[Session]
) -> None:
    with session_factory() as session:
        result = TrackerExporter(session, settings).export()
    assert result.rows == 0
    assert len(_csv_rows(settings.exports_dir / "applications.csv")) == 1
    assert load_workbook(settings.exports_dir / "applications.xlsx")["Applications"].max_row == 1


@pytest.mark.parametrize(
    "payload", ["=1+1", "+SUM(A1)", "-2+3", "@SUM(A1)", "\t=1", '=HYPERLINK("http://x","y")']
)
def test_formula_payloads_are_neutralised(
    payload: str, settings: Settings, session_factory: sessionmaker[Session]
) -> None:
    with session_factory.begin() as session:
        job = make_job("9")
        job.title = payload
        job.company = payload
        job.matching_explanation = payload
        JobRepository(session).add(job)
        application = ApplicationRepository(session).create_for_job(job)
        application.failure_reason = payload
        application.confirmation_number = payload
    with session_factory() as session:
        TrackerExporter(session, settings).export()

    for row in _csv_rows(settings.exports_dir / "applications.csv")[1:]:
        assert not any(cell.startswith(("=", "+", "-", "@", "\t")) for cell in row)
    sheet = load_workbook(settings.exports_dir / "applications.xlsx")["Applications"]
    title = sheet.cell(2, HEADERS.index("Job Title") + 1)
    assert title.data_type == "s"
    assert str(title.value).startswith("'")
    for line in sheet.iter_rows():
        assert all(cell.data_type != "f" for cell in line)


def test_clean_text_strips_control_characters_and_caps_length() -> None:
    assert clean_text("a\x00b\x07c") == "abc"
    assert len(clean_text("x" * 100_000)) <= 32_000
    assert clean_text("plain") == "plain"


def test_control_characters_do_not_break_excel_export(
    settings: Settings, session_factory: sessionmaker[Session]
) -> None:
    with session_factory.begin() as session:
        job = make_job("8")
        job.description = "ignored"
        job.title = "Engineer\x00\x08"
        JobRepository(session).add(job)
    with session_factory() as session:
        TrackerExporter(session, settings).export(("xlsx",))
    sheet = load_workbook(settings.exports_dir / "applications.xlsx")["Applications"]
    assert sheet.cell(2, HEADERS.index("Job Title") + 1).value == "Engineer"


def test_repeated_export_replaces_files_and_leaves_no_temp_files(
    settings: Settings, session_factory: sessionmaker[Session]
) -> None:
    _seed(session_factory)
    for _ in range(2):
        with session_factory() as session:
            TrackerExporter(session, settings).export()
    assert sorted(p.name for p in settings.exports_dir.iterdir()) == [
        "applications.csv",
        "applications.xlsx",
    ]
    assert len(_csv_rows(settings.exports_dir / "applications.csv")) == 3


def test_export_api_default_and_single_format(
    settings: Settings, session_factory: sessionmaker[Session]
) -> None:
    _seed(session_factory)
    with TestClient(create_app(settings)) as client:
        both = client.post("/exports/tracker")
        assert both.status_code == 200
        body = both.json()
        assert body["rows"] == 2
        assert body["files"] == {"xlsx": "applications.xlsx", "csv": "applications.csv"}
        assert str(settings.project_root) not in both.text

        (settings.exports_dir / "applications.csv").unlink()
        only_csv = client.post("/exports/tracker", json={"formats": ["csv"]})
        assert only_csv.json()["files"] == {"csv": "applications.csv"}
        assert (settings.exports_dir / "applications.csv").exists()


@pytest.mark.parametrize("formats", [[], ["pdf"], ["csv", "xlsx", "csv"]])
def test_export_api_rejects_bad_formats(settings: Settings, formats: list[str]) -> None:
    with TestClient(create_app(settings)) as client:
        assert client.post("/exports/tracker", json={"formats": formats}).status_code == 422


def test_export_api_requires_api_key(tmp_path: Path) -> None:
    keyed = Settings(
        project_root=tmp_path, config_directory=PROJECT_ROOT / "config", api_key="test-key"
    )
    with TestClient(create_app(keyed)) as client:
        assert client.post("/exports/tracker").status_code == 401
        ok = client.post("/exports/tracker", headers={"X-API-Key": "test-key"})
        assert ok.status_code == 200


def test_locked_file_returns_conflict(settings: Settings, monkeypatch: pytest.MonkeyPatch) -> None:
    def locked(src: object, dst: object) -> None:
        raise PermissionError("in use")

    monkeypatch.setattr(os, "replace", locked)
    with TestClient(create_app(settings)) as client:
        response = client.post("/exports/tracker", json={"formats": ["csv"]})
    assert response.status_code == 409
    assert "close it in Excel" in response.json()["detail"]
    assert list(settings.exports_dir.glob("*.tmp")) == []


def test_export_cli(settings: Settings, session_factory: sessionmaker[Session]) -> None:
    _seed(session_factory)
    env = {
        "PROJECT_ROOT": str(settings.project_root),
        "CONFIG_DIRECTORY": str(settings.config_dir),
    }
    runner = CliRunner()
    result = runner.invoke(cli, ["export", "--format", "csv"], env=env)
    assert result.exit_code == 0, result.output
    assert "rows=2" in result.output
    assert (settings.exports_dir / "applications.csv").exists()
    assert not (settings.exports_dir / "applications.xlsx").exists()

    bad = runner.invoke(cli, ["export", "--format", "pdf"], env=env)
    assert bad.exit_code == 1
