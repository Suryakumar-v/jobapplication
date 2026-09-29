from __future__ import annotations

import hashlib
import shutil
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
import yaml
from docx import Document
from fastapi.testclient import TestClient

from app.config import PROJECT_ROOT, Settings
from app.database.repositories import ApplicationRepository
from app.main import create_app
from app.models.application import ApplicationStatus
from app.models.candidate import CandidateProfile
from app.services.config_loader import LoadedConfig, load_matching_config
from app.services.resume_tailor import (
    build_plan,
    parse_docx,
    read_master_text,
    render_docx,
    validate_resume,
)

from .test_matching_api import import_jobs, job
from .test_matching_engine import STRONG, make_engine
from .test_matching_engine import make_job as make_posting

MASTER = Path("resumes/source/master_resume.docx")


def build_master(path: Path, profile: CandidateProfile, omit: tuple[str, ...] = ()) -> None:
    lines = [f"{profile.personal.first_name} {profile.personal.last_name}", profile.summary]
    for skill in profile.skills:
        lines.append(skill.name)
        lines.extend(skill.aliases)
    for role in profile.experience:
        lines.extend([role.employer, role.title, *role.bullets])
    lines.extend(edu.institution for edu in profile.education)
    document = Document()
    for line in lines:
        if not any(token in line for token in omit):
            document.add_paragraph(line)
    table = document.add_table(rows=1, cols=1)
    for cert in profile.certifications:
        if not any(token in cert.name for token in omit):
            table.rows[0].cells[0].add_paragraph(cert.name)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


@pytest.fixture
def real_profile_settings(tmp_path: Path) -> Settings:
    config = tmp_path / "config"
    shutil.copytree(PROJECT_ROOT / "config", config)
    shutil.copy(config / "candidate_profile.example.yaml", config / "candidate_profile.yaml")
    return Settings(project_root=tmp_path, config_directory=Path("config"))


@pytest.fixture
def profile(real_profile_settings: Settings) -> CandidateProfile:
    path = real_profile_settings.config_dir / "candidate_profile.yaml"
    return CandidateProfile.model_validate(yaml.safe_load(path.read_text(encoding="utf-8")))


@pytest.fixture
def client(real_profile_settings: Settings, profile: CandidateProfile) -> Iterator[TestClient]:
    build_master(real_profile_settings.project_root / MASTER, profile)
    with TestClient(create_app(real_profile_settings)) as test_client:
        yield test_client


def analysed_job(client: TestClient, company: str = "Acme Networks", **kwargs: Any) -> str:
    (job_id,) = import_jobs(client, job(company, **kwargs))
    assert client.post(f"/matching/analyze/{job_id}").json()["result"]["recommendation"] == "APPLY"
    return job_id


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_tailor_creates_validated_resume_and_leaves_master_untouched(
    client: TestClient, real_profile_settings: Settings, profile: CandidateProfile
) -> None:
    master = real_profile_settings.project_root / MASTER
    before = sha(master)
    job_id = analysed_job(client)
    response = client.post(f"/resumes/tailor/{job_id}")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["validation"]["passed"] is True
    assert body["validation"]["errors"] == []
    assert body["status"] == "PREPARED"
    assert body["version"] == 1
    assert body["master_resume_sha256"] == before == sha(master)
    output = real_profile_settings.project_root / body["path"]
    assert output.is_file()
    assert output != master
    assert body["path"] == f"resumes/generated/{body['application_id']}/alex-example-resume.docx"

    parsed = parse_docx(output)
    assert set(parsed.skills) == {skill.name for skill in profile.skills}
    assert parsed.skills[0] == "Network Security"
    for header, bullets in parsed.roles:
        source = next(r for r in profile.experience if r.title in header)
        assert sorted(bullets) == sorted(source.bullets)

    assert client.get(f"/jobs/{job_id}").json()["status"] == "PREPARED"
    history = client.get(f"/resumes/{body['application_id']}").json()
    assert len(history) == 1
    assert history[0]["path"] == body["path"]


def test_job_requirements_the_candidate_lacks_are_never_added(
    client: TestClient, real_profile_settings: Settings
) -> None:
    job_id = analysed_job(client, description=f"{STRONG}\nPalo Alto is a plus\nZscaler is a plus")
    body = client.post(f"/resumes/tailor/{job_id}").json()
    text = parse_docx(real_profile_settings.project_root / body["path"]).text
    assert "Palo Alto" not in text
    assert "Zscaler" not in text
    assert "Not claimed" in body["tailoring_summary"]
    assert "Palo Alto" in body["tailoring_summary"]


def test_relevant_skills_and_bullets_move_first(
    real_profile_settings: Settings, profile: CandidateProfile
) -> None:
    loaded: LoadedConfig = load_matching_config(real_profile_settings)
    engine = make_engine(loaded)
    posting = make_posting(description="Required: Python automation and Splunk")
    match = engine.evaluate(posting)
    master = "\n".join(
        [s.name for s in profile.skills]
        + [b for r in profile.experience for b in r.bullets]
        + [r.employer for r in profile.experience]
        + [r.title for r in profile.experience]
        + [e.institution for e in profile.education]
        + [c.name for c in profile.certifications]
    )
    plan = build_plan(profile, master, match, engine)
    assert set(plan.skills[:2]) == {"Splunk", "Python automation"}
    assert len(plan.skills) == len(profile.skills)
    newest = plan.roles[0]
    assert "Python" in newest.bullets[0]
    assert sorted(newest.bullets) == sorted(profile.experience[0].bullets)


def test_unanalysed_or_not_apply_jobs_are_refused(client: TestClient) -> None:
    (fresh,) = import_jobs(client, job("Fresh Co"))
    assert client.post(f"/resumes/tailor/{fresh}").status_code == 409
    (weak,) = import_jobs(client, job("Weak Co", "Required: Kubernetes, Terraform, Azure"))
    client.post(f"/matching/analyze/{weak}")
    assert client.post(f"/resumes/tailor/{weak}").status_code == 409
    assert client.post("/resumes/tailor/JOB-none").status_code == 404
    assert client.get("/resumes/APP-none").status_code == 404


def test_example_profile_is_never_used(settings: Settings) -> None:
    with TestClient(create_app(settings)) as api:
        (job_id,) = import_jobs(api, job("Acme Networks"))
        api.post(f"/matching/analyze/{job_id}")
        response = api.post(f"/resumes/tailor/{job_id}")
    assert response.status_code == 409
    assert "example" in response.json()["detail"]


def test_missing_or_unreadable_master_is_rejected(
    client: TestClient, real_profile_settings: Settings
) -> None:
    job_id = analysed_job(client)
    master = real_profile_settings.project_root / MASTER
    master.write_bytes(b"not a docx")
    assert client.post(f"/resumes/tailor/{job_id}").status_code == 422
    master.unlink()
    response = client.post(f"/resumes/tailor/{job_id}")
    assert response.status_code == 422
    assert "not found" in response.json()["detail"]
    assert client.get(f"/jobs/{job_id}").json()["status"] == "ANALYZED"


def test_master_that_contradicts_profile_blocks_resume(
    client: TestClient, real_profile_settings: Settings, profile: CandidateProfile
) -> None:
    build_master(real_profile_settings.project_root / MASTER, profile, omit=("Example Networks",))
    job_id = analysed_job(client)
    response = client.post(f"/resumes/tailor/{job_id}")
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert any("employer not found" in e for e in detail["validation"]["errors"])
    generated = real_profile_settings.generated_resumes_dir
    assert not list(generated.rglob("*.docx"))
    assert client.get(f"/jobs/{job_id}").json()["status"] == "ANALYZED"


def test_skill_missing_from_master_is_left_out_with_warning(
    client: TestClient, real_profile_settings: Settings, profile: CandidateProfile
) -> None:
    build_master(real_profile_settings.project_root / MASTER, profile, omit=("Splunk",))
    job_id = analysed_job(client)
    body = client.post(f"/resumes/tailor/{job_id}").json()
    assert body["validation"]["passed"] is True
    assert any("Splunk" in w for w in body["validation"]["warnings"])
    parsed = parse_docx(real_profile_settings.project_root / body["path"])
    assert "Splunk" not in parsed.skills


def test_regeneration_keeps_earlier_versions(
    client: TestClient, real_profile_settings: Settings
) -> None:
    job_id = analysed_job(client)
    first = client.post(f"/resumes/tailor/{job_id}").json()
    second = client.post(f"/resumes/tailor/{job_id}")
    assert second.status_code == 200
    body = second.json()
    assert body["version"] == 2
    assert body["application_id"] == first["application_id"]
    assert body["path"].endswith("alex-example-resume-v2.docx")
    root = real_profile_settings.project_root
    assert (root / first["path"]).is_file()
    assert (root / body["path"]).is_file()
    assert len(client.get(f"/resumes/{first['application_id']}").json()) == 2


def test_resume_cannot_change_after_preparation_moves_on(client: TestClient) -> None:
    job_id = analysed_job(client)
    application_id = client.post(f"/resumes/tailor/{job_id}").json()["application_id"]
    with client.app.state.session_factory() as session:  # type: ignore[attr-defined]
        ApplicationRepository(session).set_status(application_id, ApplicationStatus.FORM_STARTED)
        session.commit()
    assert client.post(f"/resumes/tailor/{job_id}").status_code == 409


def test_resume_endpoints_require_api_key(
    real_profile_settings: Settings, profile: CandidateProfile
) -> None:
    keyed = real_profile_settings.model_copy(update={"api_key": "k"})
    with TestClient(create_app(keyed)) as anonymous:
        assert anonymous.post("/resumes/tailor/JOB-1").status_code == 401
        assert anonymous.get("/resumes/APP-1").status_code == 401


@pytest.fixture
def written(
    tmp_path: Path, real_profile_settings: Settings, profile: CandidateProfile
) -> tuple[Path, str, LoadedConfig]:
    loaded = load_matching_config(real_profile_settings)
    engine = make_engine(loaded)
    match = engine.evaluate(make_posting())
    master_path = tmp_path / "master.docx"
    build_master(master_path, profile)
    master_text = read_master_text(master_path)
    output = tmp_path / "out.docx"
    render_docx(profile, build_plan(profile, master_text, match, engine), output)
    return output, master_text, loaded


def _validate(written: tuple[Path, str, LoadedConfig], path: Path | None = None) -> Any:
    output, master_text, loaded = written
    return validate_resume(
        parse_docx(path or output), loaded.candidate, master_text, make_engine(loaded)
    )


def test_untampered_document_validates(written: tuple[Path, str, LoadedConfig]) -> None:
    result = _validate(written)
    assert result.passed, result.errors
    assert {"identity", "summary", "skills", "experience", "unsupported_terms"} <= set(
        result.checks_passed
    )


def _tamper(source: Path, target: Path, kind: str) -> None:
    document = Document(str(source))
    paragraphs = document.paragraphs
    if kind == "skill":
        skills = next(p for p in paragraphs if " \u2022 " in p.text)
        skills.text = f"{skills.text} \u2022 Palo Alto"
    elif kind == "bullet":
        first = next(p for p in paragraphs if p.style is not None and p.style.name == "List Bullet")
        first.insert_paragraph_before("Led a migration to Kubernetes", style="List Bullet")
    elif kind == "certification":
        heading = next(p for p in paragraphs if p.text == "Education")
        heading.insert_paragraph_before("CISSP", style="List Bullet")
    elif kind == "education":
        document.add_paragraph("PhD \u2014 Made-up University")
    elif kind == "summary":
        summary = next(p for p in paragraphs if p.text.startswith("Network security engineer"))
        summary.text = "World-class expert in everything."
    document.save(str(target))


@pytest.mark.parametrize(
    ("kind", "expected"),
    [
        ("skill", "Skill not supported by profile and master: Palo Alto"),
        ("skill", "Mentions terms the candidate does not hold: Palo Alto"),
        ("bullet", "Bullets differ from the profile"),
        ("bullet", "Mentions terms the candidate does not hold: Kubernetes"),
        ("certification", "Unsupported certifications entry: CISSP"),
        ("education", "Unsupported education entry"),
        ("summary", "Summary differs from the profile"),
    ],
)
def test_added_claims_are_detected(
    written: tuple[Path, str, LoadedConfig], tmp_path: Path, kind: str, expected: str
) -> None:
    tampered = tmp_path / "tampered.docx"
    _tamper(written[0], tampered, kind)
    result = _validate(written, tampered)
    assert not result.passed
    assert any(expected in error for error in result.errors), result.errors
